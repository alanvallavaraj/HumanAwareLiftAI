"""Write Word copies of the title page and highlights for submission systems that require .docx.

Requires python-docx. Writes paper/title_page.docx and paper/Highlights.docx.
The highlights file carries no author metadata, so it can be sent to reviewers.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
TITLE = "Requirement-aware multi-lift dispatch: restricted waiting and accessibility trade-offs"
DOI = "https://doi.org/10.5281/zenodo.23288500"
REPO = "https://github.com/alanvallavaraj/HumanAwareLiftAI"


def base(author):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Mm(24))
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.comments = cp.keywords = cp.subject = ""
    return d


def para(d, text="", bold=False, italic=False, size=None, align=None, after=8):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


def heading(d, text):
    p = para(d, text, bold=True, size=13, after=4)
    p.paragraph_format.space_before = Pt(12)


def hyperlink(p, url):
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = url; r.append(t)
    h.append(r)
    p._p.append(h)


def title_page():
    d = base("Alan Immanuel Benjamin Vallavaraj")
    para(d, TITLE, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    p = para(d, "Alan Immanuel Benjamin Vallavaraj", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    sup = p.add_run("*"); sup.font.superscript = True
    para(d, "School of Computer Science and Engineering, University of Westminster, 115 New Cavendish Street, London W1W 6UW, United Kingdom",
         italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(d, "*Corresponding author. Email: a.vallavaraj@westminster.ac.uk", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    heading(d, "Author contributions (CRediT)")
    para(d, "Alan Immanuel Benjamin Vallavaraj: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, "
            "Data curation, Writing – original draft, Writing – review & editing, Visualization, Project administration.")
    heading(d, "Funding")
    para(d, "This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.")
    heading(d, "Declaration of competing interests")
    para(d, "The author declares no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.")
    heading(d, "Data and code availability")
    p = para(d, "All code, data and reproduction instructions are archived on Zenodo at ")
    hyperlink(p, DOI)
    p.add_run(" and mirrored at ")
    hyperlink(p, REPO)
    p.add_run(". Code is released under the MIT licence and data under the Creative Commons Attribution 4.0 licence. "
              "In the anonymised manuscript, this record is cited as the withheld dataset reference.")
    p = para(d, "Dataset reference: A. I. B. Vallavaraj, [dataset] HumanAwareLiftAI: synthetic lift-dispatch experiments and "
                "passenger-event records, Zenodo, v1.0.0 (2026). ")
    hyperlink(p, DOI)
    d.save(PAPER / "title_page.docx")


def highlights():
    bullets = [l[2:].strip() for l in (PAPER / "Highlights.txt").read_text().splitlines() if l.startswith("- ")]
    assert 3 <= len(bullets) <= 5 and all(len(b) <= 85 for b in bullets), bullets
    d = base("")
    para(d, "Highlights", bold=True, size=14, after=10)
    for b in bullets:
        p = d.add_paragraph(b, style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
    d.save(PAPER / "Highlights.docx")


if __name__ == "__main__":
    title_page()
    highlights()
    print("Wrote paper/title_page.docx and paper/Highlights.docx")
